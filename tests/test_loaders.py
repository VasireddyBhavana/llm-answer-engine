import io
import json

import pytest

from grounded.loaders import UnsupportedFile, parse_file


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal valid PDF with one text line per page (no external tools needed)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids, page_objs = [], []
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        page_objs += [f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                      f"/Contents {content_id} 0 R "
                      f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>",
                      f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"]
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    objs += page_objs + ["<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for n, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
              .encode())
    return out.getvalue()


def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph("Travel Policy", style="Title")
    d.add_heading("Booking", level=1)
    d.add_paragraph("Book flights at least 14 days in advance.")
    d.add_heading("Expenses", level=1)
    d.add_paragraph("Meals are reimbursed up to 75 dollars per day.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Hotel cap"
    t.cell(1, 0).text, t.cell(1, 1).text = "EU", "180"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_pages_become_sections():
    doc = parse_file("handbook.pdf", make_pdf(["Vacation is 20 days.", "Sick leave is 10 days."]))
    assert doc.mime == "application/pdf" and doc.meta["pages"] == 2
    assert "## Page 1" in doc.text and "## Page 2" in doc.text
    assert "Sick leave is 10 days." in doc.text


def test_docx_headings_and_tables():
    doc = parse_file("travel.docx", make_docx())
    assert doc.title == "Travel Policy"
    assert "## Booking" in doc.text and "## Expenses" in doc.text
    assert "Region: EU; Hotel cap: 180" in doc.text


def test_html_strips_scripts_and_keeps_headings():
    html = b"""<html><head><title>Status</title><script>alert('x')</script></head><body>
    <nav>menu</nav><h1>Service Status</h1><p>All systems normal.</p><h2>API</h2>
    <ul><li>Latency OK</li></ul></body></html>"""
    doc = parse_file("status.html", html)
    assert doc.title == "Status"
    assert "alert" not in doc.text and "menu" not in doc.text
    assert "# Service Status" in doc.text and "## API" in doc.text and "- Latency OK" in doc.text


def test_csv_rows_become_paragraphs():
    doc = parse_file("people.csv", b"name,team\nAda,Platform\nLin,Search\n")
    assert "name: Lin; team: Search" in doc.text and doc.meta["rows"] == 2


def test_json_is_flattened():
    doc = parse_file("cfg.json", json.dumps({"service": {"retries": 3, "hosts": ["a"]}}).encode())
    assert "service.retries: 3" in doc.text and "service.hosts[0]: a" in doc.text


def test_markdown_title_detected_and_plaintext_gets_one():
    assert parse_file("x.md", b"# Real Title\n\nBody").title == "Real Title"
    txt = parse_file("release_notes.txt", b"Fixed a bug.")
    assert txt.title == "release notes" and txt.text.startswith("# release notes")


@pytest.mark.parametrize("name,data", [("virus.exe", b"MZ"), ("empty.md", b"   \n")])
def test_rejects_unsupported_and_empty(name, data):
    with pytest.raises(UnsupportedFile):
        parse_file(name, data)
